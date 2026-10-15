"""
Document Parser and Text Extraction Service for AetherMind Module 2.
Supports PDF (PyMuPDF / pypdf), DOCX (python-docx), TXT, and CSV (pandas / standard csv).
Includes comprehensive validation for file types, file sizes, and empty or corrupt files.
"""

import io
import os
import csv
from typing import Dict, Any, Tuple, Optional

# Supported file extensions and configuration
SUPPORTED_EXTENSIONS = {
    ".pdf": "PDF Document",
    ".docx": "Word Document",
    ".txt": "Plain Text",
    ".csv": "CSV Spreadsheet",
}

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB limit

# Cache status of optional native libraries to prevent repeated slow DLL checks
_PYMUPDF_AVAILABLE: Optional[bool] = None
_PANDAS_AVAILABLE: Optional[bool] = None


class DocumentParsingError(Exception):
    """Custom exception raised when document parsing fails."""
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def validate_file_metadata(filename: str, file_size: int) -> str:
    """
    Validate file name, extension, and size.
    Returns normalized extension (e.g. '.pdf').
    """
    if not filename or not filename.strip():
        raise DocumentParsingError("File name is missing or invalid.", status_code=400)

    _, ext = os.path.splitext(filename.lower().strip())
    if not ext or ext not in SUPPORTED_EXTENSIONS:
        supported_list = ", ".join([f"{name} ({e})" for e, name in SUPPORTED_EXTENSIONS.items()])
        raise DocumentParsingError(
            f"Unsupported file format '{ext or 'unknown'}'. Supported formats: {supported_list}.",
            status_code=400,
        )

    if file_size <= 0:
        raise DocumentParsingError(
            f"The uploaded file '{filename}' is empty (0 bytes). Please upload a valid document.",
            status_code=400,
        )

    if file_size > MAX_FILE_SIZE_BYTES:
        max_mb = MAX_FILE_SIZE_BYTES // (1024 * 1024)
        actual_mb = file_size / (1024 * 1024)
        raise DocumentParsingError(
            f"File size ({actual_mb:.2f} MB) exceeds the maximum allowed limit of {max_mb} MB.",
            status_code=413,
        )

    return ext


def extract_text_from_pdf(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """
    Extract text from PDF using PyMuPDF (fitz) or pypdf.
    Gracefully handles encrypted, corrupt, or empty PDFs.
    """
    global _PYMUPDF_AVAILABLE
    
    # 1. Check if PyMuPDF can be used
    if _PYMUPDF_AVAILABLE is not False:
        try:
            import fitz
            _PYMUPDF_AVAILABLE = True
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            try:
                if doc.is_encrypted:
                    raise DocumentParsingError(
                        "The uploaded PDF is password-protected or encrypted. Please provide an unencrypted PDF.",
                        status_code=400,
                    )

                page_count = len(doc)
                if page_count == 0:
                    raise DocumentParsingError("The PDF document contains no pages.", status_code=400)

                extracted_pages = []
                total_chars = 0

                for page_index in range(page_count):
                    page = doc[page_index]
                    page_text = page.get_text().strip()
                    if page_text:
                        extracted_pages.append(f"--- Page {page_index + 1} ---\n{page_text}")
                        total_chars += len(page_text)

                full_text = "\n\n".join(extracted_pages).strip()
                if not full_text:
                    full_text = "[Notice: No selectable text could be extracted from this PDF. It may consist solely of scanned images, vector graphics, or blank pages.]"

                word_count = len(full_text.split()) if full_text else 0

                metadata = {
                    "format": "PDF",
                    "engine": "PyMuPDF",
                    "page_count": page_count,
                    "char_count": total_chars,
                    "word_count": word_count,
                    "has_selectable_text": total_chars > 0,
                }
                return full_text, metadata
            finally:
                doc.close()
        except DocumentParsingError:
            raise
        except Exception:
            _PYMUPDF_AVAILABLE = False

    # 2. Resilient pypdf extraction
    try:
        import pypdf
        stream = io.BytesIO(file_bytes)
        reader = pypdf.PdfReader(stream)
    except Exception as e:
        raise DocumentParsingError(
            f"Unable to read PDF file: The file may be corrupt or not a valid PDF document. ({str(e)})",
            status_code=422,
        )

    try:
        if reader.is_encrypted:
            raise DocumentParsingError(
                "The uploaded PDF is password-protected or encrypted. Please provide an unencrypted PDF.",
                status_code=400,
            )

        page_count = len(reader.pages)
        if page_count == 0:
            raise DocumentParsingError("The PDF document contains no pages.", status_code=400)

        extracted_pages = []
        total_chars = 0

        for page_index, page in enumerate(reader.pages):
            page_text = (page.extract_text() or "").strip()
            if page_text:
                extracted_pages.append(f"--- Page {page_index + 1} ---\n{page_text}")
                total_chars += len(page_text)

        full_text = "\n\n".join(extracted_pages).strip()
        if not full_text:
            full_text = "[Notice: No selectable text could be extracted from this PDF. It may consist solely of scanned images, vector graphics, or blank pages.]"

        word_count = len(full_text.split()) if full_text else 0

        metadata = {
            "format": "PDF",
            "engine": "pypdf",
            "page_count": page_count,
            "char_count": total_chars,
            "word_count": word_count,
            "has_selectable_text": total_chars > 0,
        }
        return full_text, metadata
    except DocumentParsingError:
        raise
    except Exception as e:
        raise DocumentParsingError(
            f"Failed to extract text from PDF: The file appears to be corrupt or malformed. ({str(e)})",
            status_code=422,
        )


def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extract text from Word Document (.docx) using python-docx."""
    try:
        import docx
        doc_stream = io.BytesIO(file_bytes)
        doc = docx.Document(doc_stream)
    except Exception as e:
        raise DocumentParsingError(
            f"Unable to read DOCX file: The file may be corrupt or not a valid Word document. ({str(e)})",
            status_code=422,
        )

    try:
        paragraphs_text = []
        for p in doc.paragraphs:
            cleaned = p.text.strip()
            if cleaned:
                paragraphs_text.append(cleaned)

        table_count = len(doc.tables)
        tables_text = []
        for t_idx, table in enumerate(doc.tables):
            rows_text = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    rows_text.append(" | ".join(row_cells))
            if rows_text:
                tables_text.append(f"--- Table {t_idx + 1} ---\n" + "\n".join(rows_text))

        all_blocks = []
        if paragraphs_text:
            all_blocks.append("\n\n".join(paragraphs_text))
        if tables_text:
            all_blocks.append("\n\n".join(tables_text))

        full_text = "\n\n".join(all_blocks).strip()

        if not full_text:
            full_text = "[Notice: The Word document contains no readable text or paragraph content.]"

        char_count = len(full_text)
        word_count = len(full_text.split()) if full_text else 0

        metadata = {
            "format": "DOCX",
            "engine": "python-docx",
            "paragraph_count": len(paragraphs_text),
            "table_count": table_count,
            "char_count": char_count,
            "word_count": word_count,
        }

        return full_text, metadata
    except Exception as e:
        raise DocumentParsingError(
            f"Failed to extract text from DOCX: The document structure may be corrupt. ({str(e)})",
            status_code=422,
        )


def extract_text_from_txt(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extract text from plain text file (.txt)."""
    decoded_text = ""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            decoded_text = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if not decoded_text:
        try:
            decoded_text = file_bytes.decode("utf-8", errors="replace")
        except Exception as e:
            raise DocumentParsingError(f"Failed to decode text file: {str(e)}", status_code=422)

    cleaned_text = decoded_text.strip()
    if not cleaned_text:
        raise DocumentParsingError("The text file is empty or contains only whitespace.", status_code=400)

    lines = cleaned_text.splitlines()
    line_count = len(lines)
    char_count = len(cleaned_text)
    word_count = len(cleaned_text.split())

    metadata = {
        "format": "TXT",
        "engine": "Standard Unicode",
        "line_count": line_count,
        "char_count": char_count,
        "word_count": word_count,
    }

    return cleaned_text, metadata


def extract_text_from_csv(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extract and format text from CSV spreadsheet using pandas or standard csv."""
    global _PANDAS_AVAILABLE
    
    # 1. Try pandas if available
    if _PANDAS_AVAILABLE is not False:
        try:
            import pandas as pd
            _PANDAS_AVAILABLE = True
            stream = io.BytesIO(file_bytes)
            try:
                df = pd.read_csv(stream, encoding="utf-8")
            except UnicodeDecodeError:
                stream.seek(0)
                df = pd.read_csv(stream, encoding="latin-1")

            row_count, col_count = df.shape
            columns_list = [str(col) for col in df.columns]

            summary_header = (
                f"--- CSV Summary ---\n"
                f"Total Rows: {row_count}\n"
                f"Total Columns: {col_count}\n"
                f"Columns: {', '.join(columns_list)}\n\n"
                f"--- Data Preview ---\n"
            )

            max_preview_rows = 100
            preview_df = df.head(max_preview_rows)
            table_str = preview_df.to_string(index=False)
            if row_count > max_preview_rows:
                table_str += f"\n\n[... and {row_count - max_preview_rows} more rows not shown in preview ...]"

            full_text = summary_header + table_str
            char_count = len(full_text)
            word_count = len(full_text.split())

            metadata = {
                "format": "CSV",
                "engine": "pandas",
                "row_count": row_count,
                "column_count": col_count,
                "columns": columns_list,
                "rows": preview_df.fillna("").astype(str).values.tolist(),
                "char_count": char_count,
                "word_count": word_count,
            }
            return full_text, metadata
        except Exception:
            _PANDAS_AVAILABLE = False

    # 2. Standard library csv reader
    try:
        decoded_content = None
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                decoded_content = file_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue

        if decoded_content is None:
            decoded_content = file_bytes.decode("utf-8", errors="replace")

        stripped_content = decoded_content.strip()
        if not stripped_content:
            raise DocumentParsingError("The CSV file is empty or contains no content.", status_code=400)

        lines = stripped_content.splitlines()
        first_line = lines[0] if lines else ""
        delimiter = ","
        if "\t" in first_line and first_line.count("\t") > first_line.count(","):
            delimiter = "\t"
        elif ";" in first_line and first_line.count(";") > first_line.count(","):
            delimiter = ";"

        reader = csv.reader(io.StringIO(stripped_content), delimiter=delimiter)
        rows = list(reader)

        if not rows or len(rows) == 0:
            raise DocumentParsingError("The CSV file contains no readable rows.", status_code=400)

        headers = [c.strip() for c in rows[0]]
        data_rows = rows[1:]
        row_count = len(data_rows)
        col_count = len(headers)

        summary_header = (
            f"--- CSV Summary ---\n"
            f"Total Rows: {row_count}\n"
            f"Total Columns: {col_count}\n"
            f"Columns: {', '.join(headers)}\n\n"
            f"--- Data Preview ---\n"
        )

        max_preview_rows = 100
        preview_rows = data_rows[:max_preview_rows]
        
        col_widths = [len(h) for h in headers]
        for r in preview_rows:
            for idx, cell in enumerate(r):
                if idx < len(col_widths):
                    col_widths[idx] = max(col_widths[idx], len(str(cell).strip()))

        header_row_str = " | ".join(h.ljust(col_widths[idx]) for idx, h in enumerate(headers))
        separator_str = "-+-".join("-" * col_widths[idx] for idx in range(len(headers)))
        
        formatted_rows = [header_row_str, separator_str]
        for r in preview_rows:
            padded_cells = [
                str(r[idx]).strip().ljust(col_widths[idx]) if idx < len(r) else "".ljust(col_widths[idx])
                for idx in range(len(headers))
            ]
            formatted_rows.append(" | ".join(padded_cells))

        if row_count > max_preview_rows:
            formatted_rows.append(f"\n[... and {row_count - max_preview_rows} more rows not shown in preview ...]")

        full_text = summary_header + "\n".join(formatted_rows)
        char_count = len(full_text)
        word_count = len(full_text.split())

        metadata = {
            "format": "CSV",
            "engine": "csv-parser",
            "row_count": row_count,
            "column_count": col_count,
            "columns": headers,
            "rows": [[str(cell).strip() for cell in r] for r in preview_rows],
            "char_count": char_count,
            "word_count": word_count,
        }

        return full_text, metadata

    except DocumentParsingError:
        raise
    except Exception as e:
        raise DocumentParsingError(
            f"Failed to parse CSV file: The file may be corrupt or malformed. ({str(e)})",
            status_code=422,
        )


def parse_document(filename: str, file_bytes: bytes) -> Dict[str, Any]:
    """
    Main entry point for document text extraction.
    Validates file metadata, routes to appropriate parser, and returns structured result.
    """
    file_size = len(file_bytes)
    ext = validate_file_metadata(filename, file_size)

    if ext == ".pdf":
        extracted_text, metadata = extract_text_from_pdf(file_bytes)
        file_type = "pdf"
    elif ext == ".docx":
        extracted_text, metadata = extract_text_from_docx(file_bytes)
        file_type = "docx"
    elif ext == ".txt":
        extracted_text, metadata = extract_text_from_txt(file_bytes)
        file_type = "txt"
    elif ext == ".csv":
        extracted_text, metadata = extract_text_from_csv(file_bytes)
        file_type = "csv"
    else:
        raise DocumentParsingError(f"Unsupported extension: {ext}", status_code=400)

    return {
        "status": "success",
        "filename": filename,
        "file_type": file_type,
        "file_size": file_size,
        "extracted_text": extracted_text,
        "metadata": metadata,
    }
